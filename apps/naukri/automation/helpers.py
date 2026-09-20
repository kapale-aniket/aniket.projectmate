import csv
import os
import random
import time
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font


def pause(page=None, min_s=1.0, max_s=2.0):
    seconds = random.uniform(min_s, max_s)
    if page is not None:
        try:
            from apps.linkedin.services.browser_preview import maybe_capture_screenshot

            maybe_capture_screenshot(page)
        except Exception:
            pass
        try:
            page.wait_for_timeout(int(seconds * 1000))
            return
        except Exception:
            pass
    time.sleep(seconds)


def dismiss_popups(page):
    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(300)
    except Exception:
        pass
    for selector in (
        'button[aria-label="Close"]',
        'button:has-text("Not Now")',
        'button:has-text("Not now")',
        'button:has-text("Skip")',
        'button:has-text("Later")',
        ".crossIcon",
    ):
        try:
            btn = page.locator(selector)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=1500)
                page.wait_for_timeout(200)
        except Exception:
            continue


def ensure_logged_in(page):
    url = (page.url or "").lower()
    if "login" in url or "nlogin" in url:
        return False
    try:
        body = (page.locator("body").inner_text() or "").lower()
    except Exception:
        body = ""
    if "login" in body and "password" in body and "mnjuser" not in url:
        return False
    return True


def write_csv(path, rows, fieldnames):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})
    return path


def write_xlsx(path, rows, fieldnames, sheet_name="Data"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for col, key in enumerate(fieldnames, start=1):
        cell = sheet.cell(row=1, column=col, value=key.replace("_", " ").title())
        cell.font = Font(bold=True)
    for row_idx, row in enumerate(rows, start=2):
        for col, key in enumerate(fieldnames, start=1):
            sheet.cell(row=row_idx, column=col, value=row.get(key, ""))
    workbook.save(path)
    return path


def today_labels():
    today = date.today()
    return {
        today.isoformat(),
        today.strftime("%d %b %Y").lower(),
        today.strftime("%d %B %Y").lower(),
        today.strftime("%d-%m-%Y"),
        today.strftime("%d/%m/%Y"),
        "today",
        datetime.now().strftime("%d %b").lower(),
    }


def text_looks_like_today(text):
    value = (text or "").strip().lower()
    if not value:
        return False
    labels = today_labels()
    if any(label in value for label in labels):
        return True
    # Relative age chips often used on Naukri
    return any(token in value for token in ("just now", "few hours", "hour ago", "hours ago"))
