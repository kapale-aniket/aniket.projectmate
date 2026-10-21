"""Read a faculty (or name) workbook into people for Perfect Connection Search."""
import re
from io import BytesIO

_TITLES = {
    "miss", "mrs", "mr", "ms", "dr", "prof", "professor",
    "shri", "smt", "kum", "ku", "adv",
}
_NAME_HEADERS = {"faculty name", "name", "full name", "person name", "faculty"}
_COLLEGE_HEADERS = {"college name", "college", "institute", "institution"}
_DISTRICT_HEADERS = {"district", "city", "location"}
_DEPT_HEADERS = {"department name", "department", "dept"}
_ROLE_HEADERS = {"designation", "role", "title"}
_UNI_HEADERS = {"university affiliation", "university", "affiliation"}


def search_name(raw):
    """Drop titles and stray digits. 'Miss BHOITE SWATI JANARDHAN' -> 'BHOITE SWATI JANARDHAN'."""
    text = str(raw or "").replace(".", " ")
    tokens = []
    for part in text.split():
        key = part.casefold().strip(".")
        if not key or key in _TITLES or key.isdigit():
            continue
        tokens.append(part)
    if not any(len(part) > 1 for part in tokens):
        return ""
    return " ".join(tokens)


def _header_map(row):
    found = {}
    for index, cell in enumerate(row):
        key = re.sub(r"\s+", " ", str(cell or "").strip()).casefold()
        if key:
            found[key] = index
    return found


def _pick(headers, names):
    for name in names:
        if name in headers:
            return headers[name]
    return None


def _cell(row, index):
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value).strip()


def people_from_workbook(source):
    """Parse xlsx bytes or a path. Sheets without a name column are skipped."""
    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        workbook = openpyxl.load_workbook(BytesIO(source), data_only=True, read_only=True)
    else:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)

    people = []
    seen = set()
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = _header_map(header_row)
            name_idx = _pick(headers, _NAME_HEADERS)
            if name_idx is None:
                continue
            college_idx = _pick(headers, _COLLEGE_HEADERS)
            district_idx = _pick(headers, _DISTRICT_HEADERS)
            dept_idx = _pick(headers, _DEPT_HEADERS)
            role_idx = _pick(headers, _ROLE_HEADERS)
            uni_idx = _pick(headers, _UNI_HEADERS)
            for row in rows:
                row = tuple(row or ())
                raw_name = _cell(row, name_idx)
                cleaned = search_name(raw_name)
                if len(cleaned.split()) < 2:
                    continue
                key = cleaned.casefold()
                if key in seen:
                    continue
                seen.add(key)
                people.append(
                    {
                        "name": cleaned,
                        "raw_name": raw_name,
                        "college": _cell(row, college_idx),
                        "district": _cell(row, district_idx),
                        "department": _cell(row, dept_idx),
                        "designation": _cell(row, role_idx),
                        "university": _cell(row, uni_idx),
                    }
                )
    finally:
        workbook.close()
    return people
