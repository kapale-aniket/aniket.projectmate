import csv
import glob
import json
import os
import re
from datetime import date, datetime
from pathlib import Path

from django.conf import settings

from apps.linkedin.models import ActivityEvent, ScraperTask, TaskLog
from apps.linkedin.services.automation_path import ensure_automation_path
from apps.linkedin.services.file_errors import friendly_file_error


def _preview_error(exc, path):
    if isinstance(exc, OSError):
        return friendly_file_error(exc, path, "preview")
    return str(exc)


def _automation_root():
    return Path(settings.LINKEDIN_AUTOMATION_ROOT)


def session_status():
    storage = _automation_root() / "common" / "storage_state.json"
    if not storage.exists():
        return {"valid": False, "message": "No session saved. Run Session Login."}
    try:
        mtime = datetime.fromtimestamp(storage.stat().st_mtime)
        with open(storage, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        cookies = data.get("cookies", [])
        has_token = any(c.get("name") == "li_at" and c.get("value") for c in cookies)
        if not has_token:
            return {"valid": False, "message": "Session file exists but login cookie missing."}
        return {
            "valid": True,
            "message": "LinkedIn session is active.",
            "updated_at": mtime.strftime("%Y-%m-%d %H:%M"),
            "path": str(storage),
        }
    except (OSError, json.JSONDecodeError) as exc:
        return {"valid": False, "message": f"Invalid session file: {exc}"}


def _jobs_from_json_data(data):
    ensure_automation_path()
    from common.tabular_io import jobs_from_json_data

    return jobs_from_json_data(data)


PREVIEW_COLUMN_ORDER = [
    "timestamp",
    "keyword",
    "location",
    "name",
    "status",
    "profile_url",
    "Profile URL",
    "job_role",
    "Job Role",
    "Job Title",
    "company_name",
    "Company Name",
    "Company",
    "location",
    "Location",
    "company_size",
    "Company Size",
    "freshers_eligible",
    "Freshers Eligible",
    "experience_required",
    "Experience Required",
    "skills_required",
    "Skills Required",
    "package_lpa",
    "Package (LPA)",
    "posted_text",
    "Posted",
    "easy_apply",
    "Easy Apply",
    "apply_link",
    "Apply Link",
    "Apply",
    "_source_file",
]


PREVIEW_HIDDEN_HEADER_KEYS = frozenset(
    {
        "srno",
        "serialno",
        "serialnumber",
        "sno",
        "slno",
        "rownumber",
        "index",
        "postedhours",
    }
)


def _normalize_header_key(header):
    return re.sub(r"[\s._#-]+", "", str(header or "").strip().lower())


def filter_preview_columns(headers):
    """Drop serial-number columns from table preview (still kept in downloaded files)."""
    return [
        h
        for h in headers
        if str(h).strip() != "#"
        and _normalize_header_key(h) not in PREVIEW_HIDDEN_HEADER_KEYS
    ]


def _order_preview_headers(headers):
    headers = filter_preview_columns(headers)
    seen = set()
    ordered = []
    for header in PREVIEW_COLUMN_ORDER:
        if header in headers and header not in seen:
            ordered.append(header)
            seen.add(header)
    for header in headers:
        if header not in seen:
            ordered.append(header)
            seen.add(header)
    return ordered


def _format_preview_cell(header, value, max_len=120):
    text = "" if value is None else str(value).strip()
    if text.startswith("http://") or text.startswith("https://"):
        return text
    return text[:max_len]


def _is_url(value):
    text = str(value or "").strip()
    return text.startswith("http://") or text.startswith("https://")


def _is_apply_header(header):
    return _normalize_header_key(header) in {"applylink", "apply"}


def _is_job_title_header(header):
    return _normalize_header_key(header) in {"jobrole", "jobtitle", "jobname"}


def _apply_url_from_record(record):
    hrefs = record.get("_hrefs") or {}
    for key, val in record.items():
        if key == "_hrefs":
            continue
        if _is_apply_header(key):
            override = hrefs.get(key)
            if override and _is_url(override):
                return str(override).strip()
            text = str(val or "").strip()
            if _is_url(text):
                return text
    for header, url in hrefs.items():
        if _is_apply_header(header) and _is_url(url):
            return str(url).strip()
    return ""


def _link_label(header, value):
    if _is_apply_header(header):
        return "Apply Now"
    if header in {"profile_url", "Profile URL"}:
        return "View profile"
    text = str(value or "").strip()
    if len(text) > 48:
        return text[:45] + "..."
    return text or "Open link"


def _preview_cell(header, record):
    raw = record.get(header, "")
    display = _format_preview_cell(header, raw)
    hrefs = record.get("_hrefs") or {}
    href = str(hrefs.get(header) or "").strip()

    if not href:
        if _is_apply_header(header):
            href = _apply_url_from_record(record)
        elif _is_job_title_header(header):
            href = _apply_url_from_record(record)
        elif _is_url(display):
            href = display

    is_link = _is_url(href)
    if is_link:
        if _is_job_title_header(header):
            label = display or "View job"
        else:
            label = _link_label(header, display or href)
        return {
            "header": header,
            "value": display,
            "is_link": True,
            "link_href": href,
            "link_label": label,
        }

    return {
        "header": header,
        "value": display or "—",
        "is_link": False,
        "link_href": "",
        "link_label": "",
    }


def _build_preview_payload(headers, records, limit=500):
    headers = _order_preview_headers(headers)
    records = records[:limit]
    rows_zipped = []
    rows = []
    for record in records:
        row_cells = []
        row_simple = []
        for header in headers:
            cell = _preview_cell(header, record)
            row_cells.append(cell)
            row_simple.append(cell["value"])
        rows_zipped.append(row_cells)
        rows.append(row_simple)
    return {"headers": headers, "rows": rows, "rows_zipped": rows_zipped, "records": records}


def apply_preview_view_filters(
    preview,
    preview_search="",
    preview_sort="",
    preview_order="asc",
    preview_page="1",
):
    """Filter, sort, and paginate preview records while keeping hyperlinks."""
    from apps.linkedin.services.filter_utils import PREVIEW_PAGE_SIZE

    if not preview or preview.get("error"):
        return preview

    headers = list(preview.get("headers") or [])
    records = list(preview.get("records") or [])

    if preview_search:
        q = preview_search.lower()
        records = [
            r
            for r in records
            if q
            in " ".join(str(v) for k, v in r.items() if k != "_hrefs").lower()
        ]

    if preview_sort and preview_sort in headers:
        reverse = preview_order == "desc"
        records = sorted(
            records,
            key=lambda r: str(r.get(preview_sort, "")).lower(),
            reverse=reverse,
        )

    page = max(1, int(preview_page or 1))
    per_page = PREVIEW_PAGE_SIZE
    total = len(records)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = min(page, total_pages)
    start = (page - 1) * per_page
    page_records = records[start : start + per_page]

    pagination = {
        "page": page,
        "per_page": per_page,
        "total_rows": total,
        "total_pages": total_pages,
        "start": start + 1 if total else 0,
        "end": min(start + per_page, total),
        "has_prev": page > 1,
        "has_next": page < total_pages,
        "previous_page": max(1, page - 1),
        "next_page": page + 1,
    }

    built = _build_preview_payload(headers, page_records)
    built.pop("records", None)
    return {**preview, **built, "pagination": pagination}


def _find_xlsx_header_row(all_rows):
    ensure_automation_path()
    from common.tabular_io import find_xlsx_header_row

    return find_xlsx_header_row(all_rows)


def _hyperlink_target(cell):
    link = cell.hyperlink
    if not link:
        return ""
    target = getattr(link, "target", None)
    if target:
        return str(target).strip()
    return str(link).strip()


def _read_xlsx_preview_records(path, limit=500):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    all_values = []
    all_hrefs = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = []
        hrefs = {}
        for col_idx, cell in enumerate(row):
            vals.append(cell.value)
            target = _hyperlink_target(cell)
            if target:
                hrefs[col_idx] = target
        all_values.append(vals)
        all_hrefs.append(hrefs)
    wb.close()

    header_idx, raw_headers = _find_xlsx_header_row(all_values)
    headers = _order_preview_headers([h for h in raw_headers if h])
    records = []
    for row_vals, href_by_col in zip(
        all_values[header_idx + 1 :], all_hrefs[header_idx + 1 :]
    ):
        if not row_vals or all(c is None or str(c).strip() == "" for c in row_vals):
            continue
        record = {}
        hrefs = {}
        for i, header in enumerate(raw_headers):
            if not header:
                continue
            value = row_vals[i] if i < len(row_vals) else ""
            record[header] = "" if value is None else value
            if i in href_by_col:
                hrefs[header] = href_by_col[i]
        if any(str(v).strip() for v in record.values()):
            if hrefs:
                record["_hrefs"] = hrefs
            records.append(record)
            if len(records) >= limit:
                break
    return headers, records


def _records_from_dicts(headers, records, limit=500):
    ordered_headers = _order_preview_headers(headers)
    normalized = []
    for record in records[:limit]:
        row = {header: record.get(header, "") for header in ordered_headers}
        normalized.append(row)
    return ordered_headers, normalized


def attach_preview_rows_zipped(preview):
    if not preview or preview.get("error"):
        return preview
    if preview.get("rows_zipped"):
        return preview
    records = preview.get("records")
    if records is not None:
        built = _build_preview_payload(preview.get("headers") or [], records)
        preview.update(built)
        return preview
    headers = preview.get("headers") or []
    rows = preview.get("rows") or []
    record_rows = [
        {header: cell for header, cell in zip(headers, row)} for row in rows
    ]
    built = _build_preview_payload(headers, record_rows)
    preview.update(built)
    return preview


def count_jobs_scraped():
    output_dir = _automation_root() / "job" / "output"
    total = 0
    for path in glob.glob(str(output_dir / "aniket_projectmate_*.json")):
        try:
            with open(path, "r", encoding="utf-8") as handle:
                data = json.load(handle)
            total += len(_jobs_from_json_data(data))
        except (OSError, json.JSONDecodeError):
            continue
    return total


def count_connections_sent():
    csv_path = _automation_root() / "connections" / "new_connection" / "connections_sent.csv"
    if not csv_path.exists():
        return 0
    try:
        with open(csv_path, "r", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
        return max(0, len(rows) - 1)
    except OSError:
        return 0


def count_posts_scraped():
    output_dir = _automation_root() / "posts" / "output"
    state_path = output_dir / "posts_state.json"
    if not state_path.exists():
        return 0
    try:
        with open(state_path, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return len(data.get("hiring_posts", []))
    except (OSError, json.JSONDecodeError):
        return 0


def count_today_records():
    today = date.today().strftime("%Y-%m-%d")
    count = 0
    job_json = _automation_root() / "job" / "output" / f"aniket_projectmate_{today}.json"
    if job_json.exists():
        try:
            with open(job_json, "r", encoding="utf-8") as handle:
                data = json.load(handle)
            count += len(_jobs_from_json_data(data))
        except (OSError, json.JSONDecodeError):
            pass
    posts_xlsx = _automation_root() / "posts" / "output" / f"hiring_posts_{today}.xlsx"
    if posts_xlsx.exists():
        count += 1
    return count


def dashboard_stats():
    return {
        "total_jobs": count_jobs_scraped(),
        "connections_sent": count_connections_sent(),
        "posts_scraped": count_posts_scraped(),
        "today_records": count_today_records(),
        "session": session_status(),
    }


def get_latest_job_output_file(prefix=None):
    """Most recent daily job Excel output for the given prefix."""
    recent = list_recent_job_output_files(limit=1, prefix=prefix)
    return recent[0] if recent else None


def list_recent_job_output_files(limit=8, prefix=None):
    """Recent job Excel outputs for quick links on the scraper page."""
    prefix = (prefix or "aniket_projectmate").strip().lower()
    type_label = "Skill Match Job" if prefix == "skill_match_job" else "Job Data"
    files = [
        f
        for f in list_output_files()
        if f.get("type") == type_label
        and str(f.get("name", "")).lower().endswith(".xlsx")
        and str(f.get("name", "")).lower().startswith(prefix)
    ]
    return files[:limit]


def resolve_output_file_path(path_str):
    """Resolve relative or absolute output path to an absolute file under automation root."""
    if not path_str or ".." in path_str:
        return None
    root = _automation_root().resolve()
    normalized = str(path_str).strip().replace("\\", "/")
    candidate = Path(path_str)
    if candidate.is_absolute():
        target = candidate.resolve()
    else:
        target = (root / normalized).resolve()
    if str(target).startswith(str(root)) and target.is_file():
        return str(target)
    for item in list_output_files():
        item_rel = str(item.get("rel_path", "")).replace("\\", "/")
        if item_rel == normalized or item.get("path") == path_str:
            return item["path"]
    return None


def output_file_key(path_str):
    """Return rel_path for URLs when possible."""
    resolved = resolve_output_file_path(path_str)
    if not resolved:
        return path_str or ""
    root = _automation_root().resolve()
    try:
        return str(Path(resolved).relative_to(root)).replace("\\", "/")
    except ValueError:
        return resolved


def list_output_files():
    files = []
    root = _automation_root()

    # Auto-purge job outputs (48h normal / 24h skill_match) whenever Output is listed.
    try:
        from apps.linkedin.services.automation_path import ensure_automation_path

        ensure_automation_path()
        from job.master_store import purge_old_files

        purge_old_files()
    except Exception:
        pass

    # Connection runs no longer produce Output files — remove leftovers.
    conn_dir = root / "connections" / "new_connection"
    if conn_dir.exists():
        for leftover in list(conn_dir.glob("*.xlsx")) + list(conn_dir.glob("*.csv")):
            try:
                leftover.unlink()
            except OSError:
                pass

    patterns = [
        (root / "job" / "output", "aniket_projectmate_*.xlsx", "Job Data"),
        (root / "job" / "output", "skill_match_job_*.xlsx", "Skill Match Job"),
        (root / "job" / "output", "merged_*.xlsx", "Job Data (Merged)"),
        (root / "posts" / "output", "hiring_posts_*.xlsx", "Posts"),
        (root / "job" / "output", "merged_*.csv", "Job Data (Merged)"),
    ]

    for directory, pattern, file_type in patterns:
        if not directory.exists():
            continue
        for path in sorted(directory.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True):
            try:
                stat = path.stat()
                record_count = _estimate_records(path)
                files.append(
                    {
                        "name": path.name,
                        "path": str(path),
                        "rel_path": str(path.relative_to(root)).replace("\\", "/"),
                        "type": file_type,
                        "date": datetime.fromtimestamp(stat.st_mtime).strftime("%b %d, %Y"),
                        "modified": datetime.fromtimestamp(stat.st_mtime).strftime("%b %d, %Y %H:%M"),
                        "mtime": stat.st_mtime,
                        "records": record_count,
                        "size_kb": round(stat.st_size / 1024, 1),
                    }
                )
            except OSError:
                continue
    return files


def _estimate_records(path):
    if path.suffix == ".json":
        try:
            with open(path, "r", encoding="utf-8") as handle:
                data = json.load(handle)
            jobs = _jobs_from_json_data(data)
            return len(jobs)
        except (OSError, json.JSONDecodeError):
            return 0
    if path.suffix == ".csv":
        try:
            with open(path, "r", encoding="utf-8") as handle:
                return max(0, sum(1 for _ in handle) - 1)
        except OSError:
            return 0
    if path.suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            ws = wb.active
            return max(0, ws.max_row - 1)
        except Exception:
            return 0
    return 0


def preview_file(path_str, limit=500):
    resolved = resolve_output_file_path(path_str)
    if not resolved:
        return {"headers": [], "rows": [], "error": "File not found"}
    path = Path(resolved)

    if path.suffix == ".json":
        try:
            with open(path, "r", encoding="utf-8") as handle:
                data = json.load(handle)
            if isinstance(data, dict) and isinstance(data.get("rows"), list):
                rows = data["rows"]
                headers = _order_preview_headers(
                    data.get("columns") or (list(rows[0].keys()) if rows else [])
                )
                headers, records = _records_from_dicts(headers, rows, limit)
                return _build_preview_payload(headers, records)
            jobs = _jobs_from_json_data(data)
            if not jobs:
                return _build_preview_payload([], [])
            headers = _order_preview_headers(list(jobs[0].keys()))
            headers, records = _records_from_dicts(headers, jobs, limit)
            return _build_preview_payload(headers, records)
        except (OSError, json.JSONDecodeError) as exc:
            return {"headers": [], "rows": [], "error": _preview_error(exc, path)}

    if path.suffix == ".csv":
        try:
            with open(path, "r", encoding="utf-8") as handle:
                reader = csv.DictReader(handle)
                headers = _order_preview_headers(list(reader.fieldnames or []))
                records = list(reader)
            headers, records = _records_from_dicts(headers, records, limit)
            return _build_preview_payload(headers, records)
        except OSError as exc:
            return {"headers": [], "rows": [], "error": _preview_error(exc, path)}

    if path.suffix == ".xlsx":
        try:
            headers, records = _read_xlsx_preview_records(path, limit)
            return _build_preview_payload(headers, records)
        except Exception as exc:
            return {"headers": [], "rows": [], "error": _preview_error(exc, path)}

    return {"headers": [], "rows": [], "error": "Unsupported file type"}


def recent_activity(limit=10):
    events = list(ActivityEvent.objects.all()[:limit])
    if events:
        return events

    fallback = []
    for task in ScraperTask.objects.exclude(status=ScraperTask.STATUS_PENDING)[:limit]:
        fallback.append(
            ActivityEvent(
                event_type=task.task_type,
                title=f"{task.get_task_type_display()} {task.status}",
                detail=json.dumps(task.config)[:200],
                created_at=task.created_at,
            )
        )
    return fallback


def collect_logs(limit=200, task_id=None, search="", sort_by="date", sort_dir="desc"):
    from apps.linkedin.services.filter_utils import sort_logs

    qs = TaskLog.objects.all()
    if task_id:
        qs = qs.filter(task_id=task_id)
    if search:
        qs = qs.filter(message__icontains=search)

    order = "-created_at" if sort_dir != "asc" else "created_at"
    if sort_by == "level":
        order = "-level" if sort_dir != "asc" else "level"
    elif sort_by == "message":
        order = "-message" if sort_dir != "asc" else "message"

    logs = list(qs.order_by(order)[:limit])
    if sort_by not in {"date", "level", "message"}:
        log_dicts = [
            {"level": log.level, "message": log.message, "created_at": log.created_at, "time": log.created_at.strftime("%H:%M:%S")}
            for log in logs
        ]
        return sort_logs(log_dicts, sort_by, sort_dir)
    return logs


def log_automation_files(limit=100):
    logs = []
    log_path = _automation_root() / "connections" / "new_connection" / "connections.log"
    if log_path.exists():
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as handle:
                for line in handle.readlines()[-limit:]:
                    logs.append({"level": "info", "message": line.rstrip(), "source": "connections"})
        except OSError:
            pass
    return logs
