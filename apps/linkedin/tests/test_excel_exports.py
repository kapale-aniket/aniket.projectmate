import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from apps.linkedin.services.excel_io import atomic_workbook_save
from apps.linkedin.services.file_merge_service import _write_xlsx


class XlsxExportTests(unittest.TestCase):
    def test_write_xlsx_creates_valid_workbook(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "merged_test.xlsx"
            headers = ["Name", "Role"]
            rows = [{"Name": "Alice", "Role": "Engineer"}, {"Name": "Bob", "Role": "Designer"}]

            _write_xlsx(path, headers, rows)

            self.assertTrue(path.exists())
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            self.assertEqual(ws.max_row, 3)
            self.assertEqual(ws.max_column, 2)
            self.assertEqual(ws["A1"].value, "Name")
            self.assertEqual(ws["A2"].value, "Alice")
            wb.close()

    def test_atomic_xlsx_write_keeps_existing_file_on_failure(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "reports.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Name", "Role"])
            ws.append(["Alice", "Engineer"])
            wb.save(path)

            original_bytes = path.read_bytes()

            def broken_save(self, target):
                Path(target).write_bytes(b"not-a-valid-xlsx")
                raise OSError("simulated write failure")

            original = Workbook.save
            Workbook.save = broken_save
            try:
                with self.assertRaises(OSError):
                    atomic_workbook_save(Workbook(), path)
            finally:
                Workbook.save = original

            self.assertEqual(path.read_bytes(), original_bytes)
