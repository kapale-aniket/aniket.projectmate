"""Shared styled Excel table rendering from format config."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common.excel_format import load_excel_format, visible_columns


def _style_dict(fmt):
    return fmt.get("style") or {}


def pattern_fill(color):
    return PatternFill("solid", fgColor=str(color).lstrip("#"))


def thin_border(color):
    side = Side(style="thin", color=str(color).lstrip("#"))
    return Border(left=side, right=side, top=side, bottom=side)


def apply_column_widths(ws, fmt):
    cols = visible_columns(fmt)
    for idx, col in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = col.get("width", 14)


def apply_watermark(ws, fmt):
    text = fmt.get("watermark_text") or ""
    if not text:
        return
    for header in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
        header.center.text = text
        header.center.size = 22
        header.center.color = "D9D9D9"
    for footer in (ws.oddFooter, ws.evenFooter, ws.firstFooter):
        footer.right.text = text
        footer.right.size = 10
        footer.right.color = "BFBFBF"


def render_title_row(ws, fmt, row=1):
    style = _style_dict(fmt)
    cols = visible_columns(fmt)
    last_col = get_column_letter(len(cols) or 1)
    ws.merge_cells(f"A{row}:{last_col}{row}")
    cell = ws.cell(row=row, column=1, value=fmt.get("sheet_title", ""))
    cell.font = Font(
        bold=True,
        size=style.get("title_font_size", 16),
        color=style.get("title_color", "1F4E79"),
    )
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = style.get("title_row_height", 36)
    freeze = style.get("freeze_panes_row", 0)
    if freeze:
        ws.freeze_panes = f"A{freeze + 1}"


def render_social_links_row(ws, fmt, row=2, social=None):
    """WhatsApp + Instagram row under the sheet title (row 2)."""
    social = social or {}
    whatsapp_url = (social.get("whatsapp_url") or "").strip()
    whatsapp_name = (social.get("whatsapp_name") or "WhatsApp").strip()
    instagram_url = (social.get("instagram_url") or "").strip()
    instagram_handle = (social.get("instagram_handle") or "Instagram").strip()
    if not whatsapp_url and not instagram_url:
        return row

    cols = visible_columns(fmt)
    col_count = len(cols) or 1
    last_col = get_column_letter(col_count)
    split = max(1, col_count // 2)
    split_col = get_column_letter(split)
    thin = Side(style="thin", color="D9D9D9")

    if whatsapp_url and instagram_url:
        ws.merge_cells(f"A{row}:{split_col}{row}")
        ws.merge_cells(
            start_row=row,
            start_column=split + 1,
            end_row=row,
            end_column=col_count,
        )
    else:
        ws.merge_cells(f"A{row}:{last_col}{row}")

    if whatsapp_url:
        wa = ws.cell(row=row, column=1, value=f"Join WhatsApp Group - {whatsapp_name}")
        wa.hyperlink = whatsapp_url
        wa.font = Font(bold=True, color="059669", underline="single", size=11)
        wa.alignment = Alignment(horizontal="center", vertical="center")
        wa_fill = pattern_fill("F0FDF4")
        end_col = split if instagram_url else col_count
        for col_idx in range(1, end_col + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = wa_fill
            cell.border = Border(
                left=Side(style="medium", color="86EFAC") if col_idx == 1 else thin,
                right=thin,
                top=thin,
                bottom=thin,
            )

    if instagram_url:
        ig_col = split + 1 if whatsapp_url else 1
        ig = ws.cell(
            row=row,
            column=ig_col,
            value=f"Follow on Instagram - {instagram_handle}",
        )
        ig.hyperlink = instagram_url
        ig.font = Font(bold=True, color="DB2777", underline="single", size=11)
        ig.alignment = Alignment(horizontal="center", vertical="center")
        ig_fill = pattern_fill("FDF2F8")
        start_col = ig_col
        for col_idx in range(start_col, col_count + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = ig_fill
            cell.border = Border(
                left=thin,
                right=Side(style="medium", color="F9A8D4") if col_idx == col_count else thin,
                top=thin,
                bottom=thin,
            )

    ws.row_dimensions[row].height = 34
    return row + 1


def render_header_row(ws, fmt, row):
    style = _style_dict(fmt)
    header_font = Font(
        bold=True,
        color=style.get("header_font_color", "FFFFFF"),
        size=style.get("header_font_size", 11),
    )
    header_fill = pattern_fill(style.get("header_bg", "1F4E79"))
    border = thin_border(style.get("border_color", "D9D9D9"))

    for col_idx, col in enumerate(visible_columns(fmt), start=1):
        cell = ws.cell(row=row, column=col_idx, value=col["header"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = style.get("header_row_height", 28)
    return row + 1


def render_data_rows(ws, fmt, row, rows, link_keys=None, center_keys=None, yes_key=None):
    style = _style_dict(fmt)
    border = thin_border(style.get("border_color", "D9D9D9"))
    alt_fill = pattern_fill(style.get("alternate_row_bg", "F2F7FB"))
    link_font = Font(color=style.get("link_color", "0563C1"), underline="single")
    yes_font = Font(bold=True, color=style.get("yes_color", "2E7D32"))
    link_keys = set(link_keys or ())
    center_keys = set(center_keys or ())
    cols = visible_columns(fmt)

    for row_idx, values in enumerate(rows, start=1):
        for col_idx, col in enumerate(cols, start=1):
            value = values.get(col["key"], "-")
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
            if col["key"] in center_keys:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if yes_key and col["key"] == yes_key and value == "Yes":
                cell.font = yes_font
            if col["key"] in link_keys and value and value != "-":
                cell.hyperlink = value
                cell.value = value
                cell.font = link_font
        ws.row_dimensions[row].height = style.get("row_height", 22)
        row += 1
    return row


def render_section_banner(ws, fmt, row, label, count_suffix=""):
    style = _style_dict(fmt)
    cols = visible_columns(fmt)
    last_col = get_column_letter(len(cols) or 1)
    ws.merge_cells(f"A{row}:{last_col}{row}")
    suffix = f"  ({count_suffix})" if count_suffix else ""
    cell = ws.cell(row=row, column=1, value=f"{label}{suffix}")
    cell.font = Font(
        bold=True,
        size=12,
        color=style.get("section_font_color", "1F4E79"),
    )
    cell.fill = pattern_fill(style.get("section_bg", "D6E4F0"))
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    return row + 1


def render_experience_banner(ws, fmt, row, label, count_suffix=""):
    style = _style_dict(fmt)
    cols = visible_columns(fmt)
    last_col = get_column_letter(len(cols) or 1)
    ws.merge_cells(f"A{row}:{last_col}{row}")
    suffix = f"  ({count_suffix})" if count_suffix else ""
    cell = ws.cell(row=row, column=1, value=f"{label}{suffix}")
    cell.font = Font(
        bold=True,
        size=13,
        color=style.get("banner_font_color", "FFFFFF"),
    )
    cell.fill = pattern_fill(style.get("banner_bg", "1F4E79"))
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30
    return row + 1


def record_row_values(idx, record, fmt):
    values = {}
    for col in visible_columns(fmt):
        key = col["key"]
        if key == "sr_no" and idx is not None:
            values[key] = idx
        else:
            values[key] = record.get(key, "-")
    return values


def write_formatted_workbook(
    path,
    fmt,
    *,
    scraper_type,
    rows,
    link_keys=None,
    center_keys=None,
    yes_key=None,
    header_row=3,
):
    from apps.linkedin.services.excel_io import atomic_workbook_save

    fmt = fmt or load_excel_format(scraper_type)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = fmt.get("sheet_name", "Sheet1")[:31]

    apply_column_widths(sheet, fmt)
    render_title_row(sheet, fmt, row=1)
    apply_watermark(sheet, fmt)

    row = render_header_row(sheet, fmt, row=header_row)
    render_data_rows(
        sheet,
        fmt,
        row,
        rows,
        link_keys=link_keys,
        center_keys=center_keys,
        yes_key=yes_key,
    )

    atomic_workbook_save(workbook, path)
    return len(rows)
