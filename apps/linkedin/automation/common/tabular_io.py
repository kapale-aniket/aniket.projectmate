"""Shared tabular file helpers for dashboard preview and merge."""

import csv
import json


def jobs_from_json_data(data):
    if isinstance(data, dict):
        jobs = data.get("jobs", [])
        return jobs if isinstance(jobs, list) else []
    if isinstance(data, list):
        return data
    return []


def normalize_header_key(key):
    return str(key or "").strip()


def find_xlsx_header_row(all_rows):
    for idx, row in enumerate(all_rows):
        cells = [normalize_header_key(c) for c in row]
        if (
            "Apply Link" in cells
            or "Job Role" in cells
            or "apply_link" in cells
            or "Job Title" in cells
            or "Apply" in cells
        ):
            return idx, cells
        if "#" in cells and "Job Title" in cells:
            return idx, cells
    if all_rows:
        return 0, [normalize_header_key(c) for c in all_rows[0]]
    return 0, []


def read_json_tabular(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    jobs = jobs_from_json_data(data)
    if not jobs:
        return [], []
    headers = []
    for job in jobs:
        for key in job.keys():
            if key not in headers:
                headers.append(key)
    rows = [{key: job.get(key, "") for key in headers} for job in jobs]
    return headers, rows


def read_csv_tabular(path):
    with open(path, "r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        rows = [dict(row) for row in reader]
    return list(headers), rows


def read_xlsx_tabular(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    all_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    header_idx, headers = find_xlsx_header_row(all_rows)
    if not headers:
        return [], []
    rows = []
    for row in all_rows[header_idx + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        record = {}
        for i, header in enumerate(headers):
            if not header:
                continue
            value = row[i] if i < len(row) else ""
            record[header] = "" if value is None else value
        if any(str(v).strip() for v in record.values()):
            rows.append(record)
    return headers, rows


def read_tabular_file(path):
    suffix = str(path).lower()
    if suffix.endswith(".json"):
        return read_json_tabular(path)
    if suffix.endswith(".csv"):
        return read_csv_tabular(path)
    if suffix.endswith(".xlsx"):
        return read_xlsx_tabular(path)
    raise ValueError(f"Unsupported file type: {path}")
