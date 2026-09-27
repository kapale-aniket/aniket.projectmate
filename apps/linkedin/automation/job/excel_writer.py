import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from common.excel_format import load_excel_format, visible_columns
from common.excel_table import (
    apply_column_widths,
    apply_watermark,
    record_row_values,
    render_data_rows,
    render_experience_banner,
    render_header_row,
    render_section_banner,
    render_social_links_row,
    render_title_row,
)
from job.job_filter import experience_bounds

NOTE_FONT = Font(italic=True, color="808080")

MIN_EMPLOYEES = 50

SIZE_GROUPS = [
    ("51-200 Employees", lambda n: n is not None and 50 < n <= 200),
    ("200+ Employees", lambda n: n is not None and n > 200),
]
OTHER_GROUP = "Other (Size Not Listed)"


def _social_links_from_config():
    try:
        from job import config as cfg
    except ImportError:
        return {}
    return {
        "whatsapp_url": (getattr(cfg, "WHATSAPP_GROUP_URL", "") or "").strip(),
        "whatsapp_name": (getattr(cfg, "WHATSAPP_GROUP_NAME", "") or "WhatsApp").strip(),
        "instagram_url": (getattr(cfg, "INSTAGRAM_URL", "") or "").strip(),
        "instagram_handle": (getattr(cfg, "INSTAGRAM_HANDLE", "") or "Instagram").strip(),
    }


def fresher_section(job):
    low, high = experience_bounds(job.get("experience_required"))

    if low is None:
        return "0"

    if low > 1 or high > 2:
        return None

    return "0" if low == 0 else "1"


def split_by_experience(jobs):
    zero_year, one_year = [], []
    for job in jobs:
        section = fresher_section(job)
        if section == "0":
            zero_year.append(job)
        elif section == "1":
            one_year.append(job)
    return zero_year, one_year


def _size_upper_bound(size):
    if not size or size == "-":
        return None
    nums = re.findall(r"\d+", size.replace(",", ""))
    if not nums:
        return None
    return max(int(n) for n in nums)


def keep_eligible_companies(jobs):
    kept = []
    for job in jobs:
        n = _size_upper_bound(job.get("company_size"))
        if n is None or n > MIN_EMPLOYEES:
            kept.append(job)
    return kept


def group_by_size(jobs):
    buckets = {label: [] for label, _ in SIZE_GROUPS}
    other = []
    for job in jobs:
        n = _size_upper_bound(job.get("company_size"))
        for label, matches in SIZE_GROUPS:
            if matches(n):
                buckets[label].append(job)
                break
        else:
            other.append(job)

    ordered = [(label, buckets[label]) for label, _ in SIZE_GROUPS]
    if other:
        ordered.append((OTHER_GROUP, other))
    return ordered


def _render_table(ws, fmt, start_row, label, jobs):
    row = render_section_banner(ws, fmt, start_row, label, f"{len(jobs)} jobs")

    if not jobs:
        cols = visible_columns(fmt)
        last_col = get_column_letter(len(cols) or 1)
        ws.merge_cells(f"A{row}:{last_col}{row}")
        note = ws.cell(row=row, column=1, value="No matching jobs in this category.")
        note.font = NOTE_FONT
        note.alignment = Alignment(horizontal="left", vertical="center")
        return row + 1

    row = render_header_row(ws, fmt, row)
    rows = [record_row_values(idx, job, fmt) for idx, job in enumerate(jobs, start=0)]
    return render_data_rows(
        ws,
        fmt,
        row,
        rows,
        link_keys={"apply_link"},
        center_keys={"sr_no", "easy_apply"},
        yes_key="easy_apply",
    )


def _render_experience_section(ws, fmt, start_row, label, jobs):
    layout = fmt.get("layout") or {}
    if layout.get("group_by_company_size", True):
        row = render_experience_banner(ws, fmt, start_row, label, f"{len(jobs)} jobs")
        for size_label, size_jobs in group_by_size(jobs):
            row = _render_table(ws, fmt, row, size_label, size_jobs)
            row += 1
        return row

    row = render_experience_banner(ws, fmt, start_row, label, f"{len(jobs)} jobs")
    row = render_header_row(ws, fmt, row)
    rows = [record_row_values(idx, job, fmt) for idx, job in enumerate(jobs, start=0)]
    return render_data_rows(
        ws,
        fmt,
        row,
        rows,
        link_keys={"apply_link"},
        center_keys={"sr_no", "easy_apply"},
        yes_key="easy_apply",
    )


def write_branded_excel(jobs, file_path, fmt=None):
    from apps.linkedin.services.excel_io import atomic_workbook_save

    fmt = fmt or load_excel_format("job")
    fmt = dict(fmt)
    try:
        from job import config as cfg

        sheet_title = (getattr(cfg, "SHEET_TITLE", "") or "").strip()
        scrape_mode = (getattr(cfg, "SCRAPE_MODE", "") or "").strip().lower()
        if scrape_mode == "perfect_match":
            sheet_title = sheet_title or "skill match job"
            fmt["sheet_name"] = "skill match job"
            fmt["sheet_title"] = sheet_title
        elif sheet_title:
            fmt["sheet_title"] = sheet_title
    except Exception:
        pass

    jobs = keep_eligible_companies(jobs)
    zero_year, one_year = split_by_experience(jobs)

    wb = Workbook()
    ws = wb.active
    ws.title = fmt.get("sheet_name", "Daily Jobs")[:31]

    apply_column_widths(ws, fmt)
    render_title_row(ws, fmt, row=1)
    next_row = render_social_links_row(ws, fmt, row=2, social=_social_links_from_config())
    apply_watermark(ws, fmt)

    layout = fmt.get("layout") or {}
    row = max(next_row, 3)
    if layout.get("group_by_experience", True):
        row = _render_experience_section(
            ws,
            fmt,
            row,
            layout.get("zero_year_section", "0 Years Experience"),
            zero_year,
        )
        row += 1
        row = _render_experience_section(
            ws,
            fmt,
            row,
            layout.get("one_year_section", "1 Year Experience"),
            one_year,
        )
    else:
        all_jobs = zero_year + one_year
        row = render_header_row(ws, fmt, row)
        rows = [record_row_values(idx, job, fmt) for idx, job in enumerate(all_jobs, start=0)]
        render_data_rows(
            ws,
            fmt,
            row,
            rows,
            link_keys={"apply_link"},
            center_keys={"sr_no", "easy_apply"},
            yes_key="easy_apply",
        )

    saved_path = atomic_workbook_save(wb, file_path)
    if saved_path and str(saved_path) != str(file_path):
        print(
            f"  Excel target was locked — saved as: {saved_path} "
            f"(close the open file to overwrite {file_path} next time)."
        )
    return len(zero_year) + len(one_year)
