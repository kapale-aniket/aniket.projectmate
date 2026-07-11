"""Seed demo LinkedIn parsing data for admin UI testing."""
from __future__ import annotations

import csv
import json
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from django.utils import timezone

from apps.linkedin.models import ActivityEvent, ScraperTask, TaskLog
from apps.linkedin.services.task_manager import ensure_settings_initialized
from apps.linkedin.services.trash_service import move_to_trash


DEMO_TAG = "linkedin_demo_seed"


class Command(BaseCommand):
    help = "Seed demo LinkedIn parsing tasks, logs, activity, and sample output files"

    def add_arguments(self, parser):
        parser.add_argument(
            "--fresh",
            action="store_true",
            help="Clear previously seeded LinkedIn demo tasks/files before recreating",
        )

    def handle(self, *args, **options):
        ensure_settings_initialized()

        root = Path(settings.LINKEDIN_AUTOMATION_ROOT)
        job_dir = root / "job" / "output"
        conn_dir = root / "connections" / "new_connection"
        posts_dir = root / "posts" / "output"
        for d in (job_dir, conn_dir, posts_dir):
            d.mkdir(parents=True, exist_ok=True)

        if options.get("fresh"):
            self._clear_demo(root)

        today = timezone.localdate()
        today_s = today.strftime("%Y-%m-%d")
        yesterday = today - timedelta(days=1)
        yesterday_s = yesterday.strftime("%Y-%m-%d")

        jobs = self._demo_jobs()
        self._write_jobs_json(job_dir / f"aniket_projectmate_{today_s}.json", jobs)
        self._write_jobs_xlsx(job_dir / f"aniket_projectmate_{today_s}.xlsx", jobs)
        self._write_jobs_json(
            job_dir / f"aniket_projectmate_{yesterday_s}.json",
            jobs[:4],
        )
        self._write_jobs_xlsx(
            job_dir / f"aniket_projectmate_{yesterday_s}.xlsx",
            jobs[:4],
        )
        # Extra file for merge testing
        self._write_jobs_xlsx(
            job_dir / f"aniket_projectmate_{today_s}_batch2.xlsx",
            jobs[4:],
        )

        connections = self._demo_connections()
        self._write_connections_csv(conn_dir / "connections_sent.csv", connections)
        self._write_connections_xlsx(conn_dir / "new_connections.xlsx", connections)

        posts = self._demo_posts()
        self._write_posts_state(posts_dir / "posts_state.json", posts)
        self._write_posts_xlsx(posts_dir / f"hiring_posts_{today_s}.xlsx", posts)

        # One disposable file for trash UI testing
        trash_demo = job_dir / "demo_trash_sample.xlsx"
        self._write_jobs_xlsx(trash_demo, jobs[:2])
        try:
            move_to_trash(str(trash_demo))
            self.stdout.write("OK Moved demo_trash_sample.xlsx into trash")
        except Exception as exc:  # noqa: BLE001
            self.stdout.write(self.style.WARNING(f"Trash seed skipped: {exc}"))

        tasks = self._seed_tasks_and_logs()
        self._seed_activity()

        self.stdout.write(self.style.SUCCESS("LinkedIn demo data seeded."))
        self.stdout.write(
            f"  Tasks: {tasks} | Jobs JSON/XLSX | Connections CSV/XLSX | Posts | Trash sample"
        )
        self.stdout.write("  Open /admin/linkedin/ to review Overview, Output, Logs.")

    def _clear_demo(self, root: Path):
        # Remove demo-tagged tasks (and cascade logs)
        deleted, _ = ScraperTask.objects.filter(config__demo_tag=DEMO_TAG).delete()
        ActivityEvent.objects.filter(detail__contains=DEMO_TAG).delete()

        patterns = [
            root / "job" / "output" / "aniket_projectmate_*.json",
            root / "job" / "output" / "aniket_projectmate_*.xlsx",
            root / "job" / "output" / "demo_trash_sample.xlsx",
            root / "connections" / "new_connection" / "connections_sent.csv",
            root / "connections" / "new_connection" / "new_connections.xlsx",
            root / "posts" / "output" / "hiring_posts_*.xlsx",
            root / "posts" / "output" / "posts_state.json",
        ]
        removed = 0
        for pattern in patterns:
            for path in Path(pattern.parent).glob(pattern.name):
                try:
                    path.unlink()
                    removed += 1
                except OSError:
                    pass
        self.stdout.write(f"OK Cleared demo tasks/events ({deleted}) and {removed} output files")

    def _demo_jobs(self):
        return [
            {
                "job_role": "Python Developer Fresher",
                "company_name": "Infosys",
                "company_size": "10001+",
                "experience_required": "0-1 years",
                "location": "Pune",
                "skills_required": "Python, Django, SQL",
                "package_lpa": "3.5-5",
                "apply_link": "https://www.linkedin.com/jobs/view/100001",
                "posted_ago": "2 hours ago",
                "keyword": "Python Developer Fresher",
            },
            {
                "job_role": "Django Developer",
                "company_name": "TCS",
                "company_size": "10001+",
                "experience_required": "0 years",
                "location": "Pune",
                "skills_required": "Django, REST, PostgreSQL",
                "package_lpa": "4-6",
                "apply_link": "https://www.linkedin.com/jobs/view/100002",
                "posted_ago": "5 hours ago",
                "keyword": "Django Developer Fresher",
            },
            {
                "job_role": "Software Engineer",
                "company_name": "Persistent Systems",
                "company_size": "5001-10000",
                "experience_required": "0-1 years",
                "location": "Pune",
                "skills_required": "Java, Spring, SQL",
                "package_lpa": "5-7",
                "apply_link": "https://www.linkedin.com/jobs/view/100003",
                "posted_ago": "1 day ago",
                "keyword": "Software Engineer",
            },
            {
                "job_role": "Data Analyst Fresher",
                "company_name": "Capgemini",
                "company_size": "10001+",
                "experience_required": "0 years",
                "location": "Bangalore",
                "skills_required": "Excel, SQL, Power BI",
                "package_lpa": "3-4.5",
                "apply_link": "https://www.linkedin.com/jobs/view/100004",
                "posted_ago": "3 hours ago",
                "keyword": "Data Analyst Fresher",
            },
            {
                "job_role": "Associate Software Engineer",
                "company_name": "Accenture",
                "company_size": "10001+",
                "experience_required": "0-1 years",
                "location": "Bangalore",
                "skills_required": "Python, Cloud basics",
                "package_lpa": "4.5-6",
                "apply_link": "https://www.linkedin.com/jobs/view/100005",
                "posted_ago": "6 hours ago",
                "keyword": "Associate Software Engineer",
            },
            {
                "job_role": "Graduate Engineer Trainee",
                "company_name": "Wipro",
                "company_size": "10001+",
                "experience_required": "0 years",
                "location": "Bangalore",
                "skills_required": "Any programming language",
                "package_lpa": "3.5-4",
                "apply_link": "https://www.linkedin.com/jobs/view/100006",
                "posted_ago": "8 hours ago",
                "keyword": "Graduate Engineer Trainee",
            },
            {
                "job_role": "Backend Developer Intern",
                "company_name": "Razorpay",
                "company_size": "1001-5000",
                "experience_required": "0 years",
                "location": "Bangalore",
                "skills_required": "Python, APIs",
                "package_lpa": "Stipend",
                "apply_link": "https://www.linkedin.com/jobs/view/100007",
                "posted_ago": "12 hours ago",
                "keyword": "Python Developer Fresher",
            },
            {
                "job_role": "Full Stack Fresher",
                "company_name": "Zoho",
                "company_size": "5001-10000",
                "experience_required": "0-1 years",
                "location": "Chennai",
                "skills_required": "JS, React, Django",
                "package_lpa": "5-8",
                "apply_link": "https://www.linkedin.com/jobs/view/100008",
                "posted_ago": "1 day ago",
                "keyword": "Software Engineer",
            },
        ]

    def _demo_connections(self):
        return [
            {
                "timestamp": timezone.now().isoformat(timespec="seconds"),
                "keyword": "Python Recruiter",
                "location": "Pune",
                "name": "Anita Sharma",
                "status": "sent",
                "profile_url": "https://www.linkedin.com/in/anita-sharma-demo",
            },
            {
                "timestamp": timezone.now().isoformat(timespec="seconds"),
                "keyword": "HR Recruiter",
                "location": "Pune",
                "name": "Rahul Mehta",
                "status": "sent",
                "profile_url": "https://www.linkedin.com/in/rahul-mehta-demo",
            },
            {
                "timestamp": timezone.now().isoformat(timespec="seconds"),
                "keyword": "Python Team Lead",
                "location": "Bangalore",
                "name": "Priya Nair",
                "status": "sent",
                "profile_url": "https://www.linkedin.com/in/priya-nair-demo",
            },
            {
                "timestamp": timezone.now().isoformat(timespec="seconds"),
                "keyword": "Talent Acquisition",
                "location": "Bangalore",
                "name": "Vikram Joshi",
                "status": "sent",
                "profile_url": "https://www.linkedin.com/in/vikram-joshi-demo",
            },
            {
                "timestamp": timezone.now().isoformat(timespec="seconds"),
                "keyword": "Python Recruiter",
                "location": "Hyderabad",
                "name": "Sneha Reddy",
                "status": "already_connected",
                "profile_url": "https://www.linkedin.com/in/sneha-reddy-demo",
            },
        ]

    def _demo_posts(self):
        return [
            {
                "author": "Tech Hiring India",
                "text": "We are hiring Python freshers in Pune. DM for JD.",
                "url": "https://www.linkedin.com/feed/update/urn:li:activity:900001",
                "posted_ago": "1h",
            },
            {
                "author": "Startup Jobs Board",
                "text": "Looking for Django developers (0-1 YOE) — Bangalore remote OK.",
                "url": "https://www.linkedin.com/feed/update/urn:li:activity:900002",
                "posted_ago": "3h",
            },
            {
                "author": "Campus Placement Hub",
                "text": "GET openings for 2025/2026 batch — apply via careers page.",
                "url": "https://www.linkedin.com/feed/update/urn:li:activity:900003",
                "posted_ago": "5h",
            },
        ]

    def _write_jobs_json(self, path: Path, jobs):
        path.write_text(
            json.dumps({"jobs": jobs, "demo_tag": DEMO_TAG}, indent=2),
            encoding="utf-8",
        )

    def _write_jobs_xlsx(self, path: Path, jobs):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Todays Jobs"
        headers = [
            "Sr No",
            "Job Role",
            "Company Name",
            "Employee Size",
            "Experience Required",
            "Location",
            "Skills Required",
            "Package (LPA)",
            "Apply Link",
            "Posted",
            "Keyword",
        ]
        ws.append(headers)
        for i, job in enumerate(jobs, start=1):
            ws.append(
                [
                    i,
                    job.get("job_role", ""),
                    job.get("company_name", ""),
                    job.get("company_size", ""),
                    job.get("experience_required", ""),
                    job.get("location", ""),
                    job.get("skills_required", ""),
                    job.get("package_lpa", ""),
                    job.get("apply_link", ""),
                    job.get("posted_ago", ""),
                    job.get("keyword", ""),
                ]
            )
        wb.save(path)

    def _write_connections_csv(self, path: Path, rows):
        headers = ["timestamp", "keyword", "location", "name", "status", "profile_url"]
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)

    def _write_connections_xlsx(self, path: Path, rows):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Connections"
        headers = ["Timestamp", "Keyword", "Location", "Name", "Status", "Profile URL"]
        ws.append(headers)
        for row in rows:
            ws.append(
                [
                    row.get("timestamp", ""),
                    row.get("keyword", ""),
                    row.get("location", ""),
                    row.get("name", ""),
                    row.get("status", ""),
                    row.get("profile_url", ""),
                ]
            )
        wb.save(path)

    def _write_posts_state(self, path: Path, posts):
        path.write_text(
            json.dumps({"hiring_posts": posts, "demo_tag": DEMO_TAG}, indent=2),
            encoding="utf-8",
        )

    def _write_posts_xlsx(self, path: Path, posts):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Hiring Posts"
        ws.append(["Author", "Post Text", "URL", "Posted"])
        for post in posts:
            ws.append(
                [
                    post.get("author", ""),
                    post.get("text", ""),
                    post.get("url", ""),
                    post.get("posted_ago", ""),
                ]
            )
        wb.save(path)

    def _seed_tasks_and_logs(self):
        now = timezone.now()
        specs = [
            (
                ScraperTask.TASK_SESSION,
                ScraperTask.STATUS_COMPLETED,
                {},
                [
                    (TaskLog.LEVEL_INFO, "Opened Chromium for LinkedIn login"),
                    (TaskLog.LEVEL_SUCCESS, "Session cookies saved (demo)"),
                ],
                now - timedelta(hours=5),
                now - timedelta(hours=4, minutes=58),
            ),
            (
                ScraperTask.TASK_JOB,
                ScraperTask.STATUS_COMPLETED,
                {
                    "keywords": ["Python Developer Fresher", "Django Developer Fresher"],
                    "locations": ["Pune", "Bangalore"],
                    "max_pages": 5,
                    "max_jobs": 20,
                },
                [
                    (TaskLog.LEVEL_INFO, "Starting job scraper (demo)"),
                    (TaskLog.LEVEL_INFO, "Searching: Python Developer Fresher × Pune"),
                    (TaskLog.LEVEL_INFO, "Found 12 jobs on page 1"),
                    (TaskLog.LEVEL_INFO, "Searching: Django Developer Fresher × Bangalore"),
                    (TaskLog.LEVEL_SUCCESS, "Wrote aniket_projectmate demo Excel"),
                ],
                now - timedelta(hours=3),
                now - timedelta(hours=2, minutes=40),
            ),
            (
                ScraperTask.TASK_CONNECTION,
                ScraperTask.STATUS_COMPLETED,
                {
                    "keywords": ["Python Recruiter", "HR Recruiter"],
                    "locations": ["Pune"],
                    "max_connections_per_run": 10,
                },
                [
                    (TaskLog.LEVEL_INFO, "Connection scraper started (demo)"),
                    (TaskLog.LEVEL_INFO, "Sent invite to Anita Sharma"),
                    (TaskLog.LEVEL_INFO, "Sent invite to Rahul Mehta"),
                    (TaskLog.LEVEL_SUCCESS, "5 connection rows written"),
                ],
                now - timedelta(hours=2),
                now - timedelta(hours=1, minutes=45),
            ),
            (
                ScraperTask.TASK_POSTS,
                ScraperTask.STATUS_COMPLETED,
                {"max_scroll_rounds": 10, "max_posts_per_run": 20},
                [
                    (TaskLog.LEVEL_INFO, "Scanning feed for hiring posts (demo)"),
                    (TaskLog.LEVEL_INFO, "Matched hiring keywords in 3 posts"),
                    (TaskLog.LEVEL_SUCCESS, "Posts Excel saved"),
                ],
                now - timedelta(hours=1),
                now - timedelta(minutes=50),
            ),
            (
                ScraperTask.TASK_JOB,
                ScraperTask.STATUS_STOPPED,
                {"keywords": ["Java Developer Fresher"], "locations": ["Pune"]},
                [
                    (TaskLog.LEVEL_INFO, "Job scraper started"),
                    (TaskLog.LEVEL_WARNING, "Stopped by user (demo)"),
                ],
                now - timedelta(minutes=40),
                now - timedelta(minutes=35),
            ),
            (
                ScraperTask.TASK_CONNECTION,
                ScraperTask.STATUS_FAILED,
                {"keywords": ["Recruiter"], "locations": ["Mumbai"]},
                [
                    (TaskLog.LEVEL_INFO, "Connection scraper started"),
                    (TaskLog.LEVEL_ERROR, "Demo failure: session expired — re-run Session login"),
                ],
                now - timedelta(minutes=25),
                now - timedelta(minutes=24),
            ),
        ]

        created = 0
        for task_type, status, config, logs, started, finished in specs:
            config = {**config, "demo_tag": DEMO_TAG}
            task = ScraperTask.objects.create(
                task_type=task_type,
                status=status,
                config=config,
                progress={
                    "percent": 100 if status == ScraperTask.STATUS_COMPLETED else 40,
                    "message": "Demo seed run",
                    "jobs_found": 8 if task_type == ScraperTask.TASK_JOB else 0,
                    "connections_sent": 5 if task_type == ScraperTask.TASK_CONNECTION else 0,
                    "posts_checked": 3 if task_type == ScraperTask.TASK_POSTS else 0,
                },
                error_message=(
                    "Demo failure: session expired — re-run Session login"
                    if status == ScraperTask.STATUS_FAILED
                    else ""
                ),
                run_in_background=True,
            )
            # created_at is auto; update timing fields
            ScraperTask.objects.filter(pk=task.pk).update(
                started_at=started,
                finished_at=finished,
            )
            for level, message in logs:
                log = TaskLog.objects.create(task=task, level=level, message=message)
                TaskLog.objects.filter(pk=log.pk).update(
                    created_at=started + timedelta(seconds=10 * (logs.index((level, message)) + 1))
                )
            created += 1
        return created

    def _seed_activity(self):
        events = [
            ("session", "LinkedIn session saved", "Demo session login completed"),
            ("job", "Job scrape finished", "8 demo jobs written to Excel"),
            ("connection", "Connections sent", "5 demo connection rows"),
            ("posts", "Hiring posts scraped", "3 demo posts saved"),
            ("job", "Job scrape stopped", "User stopped demo run"),
        ]
        for event_type, title, detail in events:
            ActivityEvent.objects.create(
                event_type=event_type,
                title=title,
                detail=f"{detail} [{DEMO_TAG}]",
            )
